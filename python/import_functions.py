"""
import_functions.py
Port terus dari process.php - function sama, flow sama:
  normalize_header, get_row_value, parse_date, parse_amount,
  get_company, load_file, process_file

VERSI OPTIMIZED untuk fail besar (contoh 1 juta row):
  - Company di-cache dalam memory (elak query berulang)
  - Insert guna executemany() dalam batch, bukan satu-satu
  - Untuk tax invoice, order_id -> order_db_id di-pre-fetch sekali
    guna satu query IN (...), bukan SELECT per row
  - Commit setiap N row (chunked), bukan satu transaction gergasi

Install dulu (guna --break-system-packages kalau perlu):
  pip install pymysql openpyxl --break-system-packages
"""

import csv
import hashlib
import io
import re
from datetime import datetime, timedelta

import openpyxl


# Berapa row nak proses sebelum flush batch ke DB + commit.
# Boleh tune ikut keperluan (500-2000 biasanya sweet spot).
BATCH_SIZE = 1000


# ============================================================
# HELPER FUNCTIONS (tak berubah)
# ============================================================

def normalize_header(header):
    """Sama macam normalizeHeader() dalam PHP."""
    if header is None:
        return ""
    header = str(header).strip()
    header = header.replace("\ufeff", "")  # buang BOM kalau ada
    header = header.lower()
    header = re.sub(r"[^a-z0-9]", "", header)
    return header


def get_row_value(row, header_map, keys):
    """Sama macam getRowValue() dalam PHP."""
    for key in keys:
        normalized = normalize_header(key)
        if normalized in header_map:
            idx = header_map[normalized]
            if idx < len(row):
                return row[idx]
            return ""
    return ""


def parse_date(value):
    """Sama macam parseDate() dalam PHP."""
    if value is None:
        return None
    if isinstance(value, datetime):
        return value.strftime("%Y-%m-%d %H:%M:%S")

    value = str(value).strip()
    if value == "":
        return None

    # Excel serial date number
    try:
        numeric_value = float(value)
        if 40000 < numeric_value < 60000:
            base = datetime(1899, 12, 30)
            dt = base + timedelta(days=numeric_value)
            return dt.strftime("%Y-%m-%d %H:%M:%S")
    except ValueError:
        pass

    formats = [
        "%d/%m/%Y %H:%M:%S", "%d/%m/%Y %H:%M", "%d/%m/%Y",
        "%Y-%m-%d %H:%M:%S", "%Y-%m-%d %H:%M", "%Y-%m-%d",
        "%d-%m-%Y %H:%M:%S", "%d-%m-%Y %H:%M", "%d-%m-%Y",
    ]
    for fmt in formats:
        try:
            dt = datetime.strptime(value, fmt)
            return dt.strftime("%Y-%m-%d %H:%M:%S")
        except ValueError:
            continue

    return None


def parse_amount(value):
    """Sama macam parseAmount() dalam PHP."""
    if value is None:
        return 0.00
    text = str(value).strip()
    if text == "":
        return 0.00
    try:
        return float(text)
    except ValueError:
        pass
    for token in ["RM", "SGD", "USD", ",", " ", "(", ")", '"']:
        text = text.replace(token, "")
    try:
        return float(text)
    except ValueError:
        return 0.00


def get_order_prefix(order_id):
    """Tentukan prefix dari order_id sahaja (tiada DB call).
    Dipisahkan dari get_company() supaya caching senang dibuat.
    """
    order_id = (order_id or "").strip().upper()
    if order_id.startswith("MYHQ"):
        return "MYHQ"
    elif order_id.startswith("MYBTCSB") or order_id.startswith("MYBT"):
        return "MYBT"
    elif order_id.startswith("SGHQ") or order_id.startswith("SG"):
        return "SGHQ"
    elif order_id.startswith("BN"):
        return "BNHQ"
    return "MYHQ"  # default


# ============================================================
# COMPANY DETECTION (dengan caching)
# ============================================================

_COMPANY_MAP = {
    "MYHQ": {"code": "MY", "currency": "MYR"},
    "MYBT": {"code": "MY", "currency": "MYR"},
    "SGHQ": {"code": "SG", "currency": "SGD"},
    "BNHQ": {"code": "BN", "currency": "BND"},
}


def get_company(order_id, conn):
    """Sama macam getCompany() asal - kekal untuk backward compatibility
    (contoh dipanggil sekali untuk detect company fail secara keseluruhan).
    Untuk loop utama proses row, guna get_company_cached() supaya
    tak query DB setiap row.
    """
    prefix = get_order_prefix(order_id)
    company_code = _COMPANY_MAP.get(prefix, {}).get("code", "MY")
    currency = _COMPANY_MAP.get(prefix, {}).get("currency", "MYR")

    with conn.cursor() as cur:
        cur.execute(
            "SELECT id, invoice_prefix FROM companies WHERE company_code=%s AND is_active=1 LIMIT 1",
            (company_code,),
        )
        result = cur.fetchone()

        if not result:
            cur.execute(
                "INSERT INTO companies (company_code, company_name, currency_code, invoice_prefix) "
                "VALUES (%s, %s, %s, %s)",
                (company_code, "Malaysia" if company_code == "MY" else "Singapore", currency, prefix),
            )
            company_id = cur.lastrowid
            invoice_prefixes = prefix
        else:
            company_id = result["id"]
            invoice_prefixes = result["invoice_prefix"] or ""
            prefixes = [p.strip() for p in invoice_prefixes.split(",") if p.strip()]
            if prefix not in prefixes:
                prefixes.append(prefix)
                cur.execute(
                    "UPDATE companies SET invoice_prefix=%s WHERE id=%s",
                    (",".join(prefixes), company_id),
                )

    return {
        "company_id": company_id,
        "company_code": company_code,
        "currency_code": currency,
        "invoice_prefix": prefix,
    }


class CompanyCache:
    """Cache get_company() result ikut prefix, supaya untuk 1 juta row
    yang mungkin cuma ada 2-4 prefix berbeza, DB cuma di-hit
    sekali setiap prefix (bukan sekali setiap row).
    """

    def __init__(self, conn):
        self.conn = conn
        self._cache = {}  # prefix -> company dict

    def get(self, order_id):
        prefix = get_order_prefix(order_id)
        if prefix not in self._cache:
            self._cache[prefix] = get_company(order_id, self.conn)
        return self._cache[prefix]


# ============================================================
# LOAD FILE
# ============================================================

def load_file(file_bytes, original_name, delimiter):
    """Sama macam loadFile() dalam PHP. Return list of rows (list of list).

    NOTA untuk fail sangat besar (1 juta+ row): function ni tetap
    load semua row ke memory sebagai list Python. Untuk kebanyakan
    kes (1 juta row x 30 column) ni biasanya okay (~1-2GB RAM) tapi
    kalau server RAM terhad, function ni adalah candidate seterusnya
    untuk ditukar kepada generator/streaming (proses chunk-by-chunk
    terus dari openpyxl/csv reader tanpa simpan semua dalam list).
    """
    ext = original_name.rsplit(".", 1)[-1].lower() if "." in original_name else ""
    if ext not in ("xlsx", "csv"):
        raise ValueError("Only .xlsx and .csv files are supported.")

    rows = []
    if ext == "csv":
        text = file_bytes.decode("utf-8-sig", errors="replace")
        reader = csv.reader(io.StringIO(text), delimiter=(delimiter or ","))
        for r in reader:
            rows.append(r)
    else:
        wb = openpyxl.load_workbook(io.BytesIO(file_bytes), data_only=True, read_only=True)
        ws = wb.active
        for r in ws.iter_rows(values_only=True):
            rows.append(list(r))

    if len(rows) < 2:
        raise ValueError("File is empty.")
    return rows


# ============================================================
# PROCESS FILE  (== processFile() dalam PHP, versi optimized)
# ============================================================

ORDER_SQL = """
    INSERT INTO orders (
        company_id, import_batch_id, order_id, order_datetime, member_code, member_type,
        member_name, remark, delivery_method, mobile_no, order_type, sub_total, shipping_fee,
        voucher_discount, discount, convenience_fee, order_total, gst, total_bv, total_pv,
        total_tp, payment_mode, order_status, delivery_status, payment_gateway,
        payment_gateway_id, currency_code, invoice_prefix
    ) VALUES (%s,%s,%s,%s,%s,%s,%s,%s,%s,%s,%s,%s,%s,%s,%s,%s,%s,%s,%s,%s,%s,%s,%s,%s,%s,%s,%s,%s)
    ON DUPLICATE KEY UPDATE
        import_batch_id=VALUES(import_batch_id), order_datetime=VALUES(order_datetime),
        member_code=VALUES(member_code), member_type=VALUES(member_type),
        member_name=VALUES(member_name), remark=VALUES(remark),
        delivery_method=VALUES(delivery_method), mobile_no=VALUES(mobile_no),
        order_type=VALUES(order_type), sub_total=VALUES(sub_total),
        shipping_fee=VALUES(shipping_fee), voucher_discount=VALUES(voucher_discount),
        discount=VALUES(discount), convenience_fee=VALUES(convenience_fee),
        order_total=VALUES(order_total), gst=VALUES(gst), total_bv=VALUES(total_bv),
        total_pv=VALUES(total_pv), total_tp=VALUES(total_tp), payment_mode=VALUES(payment_mode),
        order_status=VALUES(order_status), delivery_status=VALUES(delivery_status),
        payment_gateway=VALUES(payment_gateway), payment_gateway_id=VALUES(payment_gateway_id),
        currency_code=VALUES(currency_code), invoice_prefix=VALUES(invoice_prefix)
"""

ITEM_SQL = """
    INSERT INTO order_items (
        order_id, commission_month, cdo, cdo_created_date, product_type,
        item_code, item_description, brand, email, total_weight, qty,
        bv, pv, total_bv, total_pv, total_retail_price, discount,
        invoice_amount, total_invoice_amount_paid, order_processed_location
    ) VALUES (%s,%s,%s,%s,%s,%s,%s,%s,%s,%s,%s,%s,%s,%s,%s,%s,%s,%s,%s,%s)
"""


def _detect_file_type(header_map):
    has_order_id = "orderid" in header_map
    has_datetime = ("datetime" in header_map or "orderdatetime" in header_map
                     or "dateandtime" in header_map)
    has_commission = "commissionmonth" in header_map
    has_item_code = ("itemcode" in header_map or "sku" in header_map or "itemno" in header_map)

    is_order_history = has_order_id and has_datetime
    is_tax_invoice = has_commission and has_item_code

    if not is_order_history and not is_tax_invoice:
        order_history_cols = ["memberid", "membertype", "membername", "subtotal", "ordertotal"]
        has_order_cols = all(c in header_map for c in order_history_cols)
        if has_order_cols and has_order_id:
            is_order_history = True

        tax_cols = ["itemdescription", "qty", "invoiceamount"]
        has_tax_cols = all(c in header_map for c in tax_cols)
        if has_tax_cols and has_commission and has_item_code:
            is_tax_invoice = True

    return is_order_history, is_tax_invoice


def _prefetch_order_ids(conn, company_id, order_ids):
    """Satu query IN (...) untuk resolve order_id -> order_db_id,
    dipanggil sekali sebelum loop tax invoice (bukan per-row).
    MySQL ada had bilangan parameter dalam IN(), jadi kita chunk
    order_ids kepada kumpulan kecil sebelum query.
    """
    mapping = {}
    order_ids = list(order_ids)
    chunk = 2000
    with conn.cursor() as cur:
        for i in range(0, len(order_ids), chunk):
            batch = order_ids[i:i + chunk]
            placeholders = ",".join(["%s"] * len(batch))
            cur.execute(
                f"SELECT id, order_id FROM orders WHERE company_id=%s AND order_id IN ({placeholders})",
                (company_id, *batch),
            )
            for r in cur.fetchall():
                mapping[r["order_id"]] = r["id"]
    return mapping


def process_file(file_bytes, original_name, conn, delimiter=","):
    file_hash = hashlib.sha256(file_bytes).hexdigest()

    with conn.cursor() as cur:
        cur.execute(
            "SELECT id, status FROM import_batches WHERE file_hash=%s LIMIT 1",
            (file_hash,),
        )
        existing = cur.fetchone()

    if existing and existing["status"] == "completed":
        return {"success": 0, "skipped": True, "message": "File already imported."}

    if existing:
        with conn.cursor() as cur:
            cur.execute("UPDATE import_batches SET file_hash=NULL WHERE id=%s", (existing["id"],))
        conn.commit()

    rows = load_file(file_bytes, original_name, delimiter)

    # Build header map
    raw_headers = rows[0]
    header_map = {}
    for i, header in enumerate(raw_headers):
        header_map[normalize_header(header)] = i

    is_order_history, is_tax_invoice = _detect_file_type(header_map)

    if not is_order_history and not is_tax_invoice:
        raise ValueError("Unknown file format. Headers: " + ", ".join(str(h) for h in raw_headers))

    # Determine company from first order id found (untuk batch record)
    first_order_id = ""
    for row in rows[1:]:
        val = str(get_row_value(row, header_map, ["orderid", "orderno"]) or "").strip()
        if val:
            first_order_id = val
            break

    company_cache = CompanyCache(conn)
    company = company_cache.get(first_order_id)
    conn.commit()

    # Create batch
    with conn.cursor() as cur:
        cur.execute(
            "INSERT INTO import_batches (company_id, file_type, original_filename, file_hash, total_rows, status) "
            "VALUES (%s, %s, %s, %s, %s, %s)",
            (
                company["company_id"],
                "ORDER_HISTORY" if is_order_history else "TAX_INVOICE",
                original_name,
                file_hash,
                len(rows) - 1,
                "processing",
            ),
        )
        batch_id = cur.lastrowid
    conn.commit()

    success = 0
    failed = 0

    try:
        if is_order_history:
            success, failed = _process_order_history(
                conn, rows, header_map, company_cache, batch_id
            )
        else:
            success, failed = _process_tax_invoice(
                conn, rows, header_map, company_cache
            )

        status = "completed_with_errors" if failed else "completed"
        with conn.cursor() as cur:
            cur.execute(
                "UPDATE import_batches SET successful_rows=%s, failed_rows=%s, status=%s, imported_at=NOW() WHERE id=%s",
                (success, failed, status, batch_id),
            )
        conn.commit()

    except Exception as e:
        conn.rollback()
        with conn.cursor() as cur:
            cur.execute(
                "UPDATE import_batches SET status='failed', error_message=%s WHERE id=%s",
                (str(e), batch_id),
            )
        conn.commit()
        raise

    return {"success": success, "failed": failed}


def _process_order_history(conn, rows, header_map, company_cache, batch_id):
    """Batch insert untuk order history. Commit setiap BATCH_SIZE row."""
    success = 0
    failed = 0
    batch_params = []

    with conn.cursor() as cur:
        for row in rows[1:]:
            has_data = any(str(v).strip() != "" for v in row if v is not None)
            if not has_data:
                continue

            order_id = str(get_row_value(row, header_map, ["orderid", "orderno"]) or "").strip()
            if not order_id:
                failed += 1
                continue

            comp = company_cache.get(order_id)

            dt = parse_date(get_row_value(
                row, header_map, ["dateandtime", "orderdatetime", "datetime", "orderdate"]))
            if not dt:
                failed += 1
                continue

            batch_params.append((
                comp["company_id"], batch_id, order_id, dt,
                get_row_value(row, header_map, ["memberid"]),
                get_row_value(row, header_map, ["membertype"]),
                get_row_value(row, header_map, ["membername"]),
                get_row_value(row, header_map, ["remark"]),
                get_row_value(row, header_map, ["deliverymethod"]),
                get_row_value(row, header_map, ["mobileno"]),
                get_row_value(row, header_map, ["ordertype"]),
                parse_amount(get_row_value(row, header_map, ["subtotal"])),
                parse_amount(get_row_value(row, header_map, ["shippingfee"])),
                parse_amount(get_row_value(row, header_map, ["voucherdiscount"])),
                parse_amount(get_row_value(row, header_map, ["discount"])),
                parse_amount(get_row_value(row, header_map, ["conveniencefee"])),
                parse_amount(get_row_value(row, header_map, ["ordertotal"])),
                parse_amount(get_row_value(row, header_map, ["gst"])),
                parse_amount(get_row_value(row, header_map, ["totalbv"])),
                parse_amount(get_row_value(row, header_map, ["totalpv"])),
                parse_amount(get_row_value(row, header_map, ["totaltp"])),
                get_row_value(row, header_map, ["paymentmode"]),
                get_row_value(row, header_map, ["orderstatus"]),
                get_row_value(row, header_map, ["deliverystatus"]),
                get_row_value(row, header_map, ["paymentgateway"]),
                get_row_value(row, header_map, ["paymentgatewayid"]),
                comp["currency_code"],
                comp["invoice_prefix"],
            ))
            success += 1

            if len(batch_params) >= BATCH_SIZE:
                cur.executemany(ORDER_SQL, batch_params)
                conn.commit()
                batch_params = []

        if batch_params:
            cur.executemany(ORDER_SQL, batch_params)
            conn.commit()

    return success, failed


def _process_tax_invoice(conn, rows, header_map, company_cache):
    """Batch insert untuk tax invoice.
    Pre-fetch semua order_id -> order_db_id sekali (per company),
    dan pre-delete existing order_items untuk order yang terlibat,
    supaya tak perlu SELECT + DELETE per row dalam loop.
    """
    success = 0
    failed = 0

    # 1) Kumpul semua order_id unique dalam fail, ikut company
    order_ids_by_company = {}  # company_id -> set(order_id)
    order_id_to_company = {}
    for row in rows[1:]:
        has_data = any(str(v).strip() != "" for v in row if v is not None)
        if not has_data:
            continue
        order_id = str(get_row_value(row, header_map, ["orderid", "orderno"]) or "").strip()
        if not order_id:
            continue
        comp = company_cache.get(order_id)
        order_ids_by_company.setdefault(comp["company_id"], set()).add(order_id)
        order_id_to_company[order_id] = comp["company_id"]

    # 2) Pre-fetch order_db_id untuk semua order_id sekaligus (per company)
    order_lookup = {}  # (company_id, order_id) -> order_db_id
    for company_id, order_ids in order_ids_by_company.items():
        mapping = _prefetch_order_ids(conn, company_id, order_ids)
        for oid, db_id in mapping.items():
            order_lookup[(company_id, oid)] = db_id

    # 3) Pre-delete existing items untuk semua order yang bakal di-replace,
    #    guna satu DELETE ... WHERE order_id IN (...) chunked.
    all_order_db_ids = list(order_lookup.values())
    with conn.cursor() as cur:
        chunk = 2000
        for i in range(0, len(all_order_db_ids), chunk):
            batch = all_order_db_ids[i:i + chunk]
            if not batch:
                continue
            placeholders = ",".join(["%s"] * len(batch))
            cur.execute(f"DELETE FROM order_items WHERE order_id IN ({placeholders})", batch)
    conn.commit()

    # 4) Loop row, batch insert item
    batch_params = []
    with conn.cursor() as cur:
        for row in rows[1:]:
            has_data = any(str(v).strip() != "" for v in row if v is not None)
            if not has_data:
                continue

            order_id = str(get_row_value(row, header_map, ["orderid", "orderno"]) or "").strip()
            if not order_id:
                failed += 1
                continue

            company_id = order_id_to_company.get(order_id)
            order_db_id = order_lookup.get((company_id, order_id))
            if not order_db_id:
                failed += 1
                continue

            item_code = str(get_row_value(row, header_map, ["itemcode", "sku", "itemno"]) or "").strip()
            if not item_code:
                failed += 1
                continue

            batch_params.append((
                order_db_id,
                get_row_value(row, header_map, ["commissionmonth"]),
                get_row_value(row, header_map, ["cdo"]),
                parse_date(get_row_value(row, header_map, ["cdocreateddate"])),
                get_row_value(row, header_map, ["producttype"]),
                item_code,
                get_row_value(row, header_map, ["itemdescription"]),
                get_row_value(row, header_map, ["brand"]),
                get_row_value(row, header_map, ["email"]),
                parse_amount(get_row_value(row, header_map, ["totalweight"])),
                int(parse_amount(get_row_value(row, header_map, ["qty"]))),
                parse_amount(get_row_value(row, header_map, ["bv"])),
                parse_amount(get_row_value(row, header_map, ["pv"])),
                parse_amount(get_row_value(row, header_map, ["totalbv"])),
                parse_amount(get_row_value(row, header_map, ["totalpv"])),
                parse_amount(get_row_value(row, header_map, ["totalretailprice"])),
                parse_amount(get_row_value(row, header_map, ["discount"])),
                parse_amount(get_row_value(row, header_map, ["invoiceamount"])),
                parse_amount(get_row_value(row, header_map, ["totalinvoiceamountpaid"])),
                get_row_value(row, header_map, ["orderprocessedlocation"]),
            ))
            success += 1

            if len(batch_params) >= BATCH_SIZE:
                cur.executemany(ITEM_SQL, batch_params)
                conn.commit()
                batch_params = []

        if batch_params:
            cur.executemany(ITEM_SQL, batch_params)
            conn.commit()

    return success, failed